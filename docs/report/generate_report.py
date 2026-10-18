from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ROW_HEIGHT_RULE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.shared import Twips


OUT = "docs/report/УСР1_СБД_VibeCity.docx"


# ---------------------------------------------------------------------------
# Вспомогательные функции оформления
# ---------------------------------------------------------------------------

def set_run_font(run, size=None, bold=None, color=None):
    run.font.name = "Times New Roman"
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        rfonts.set(qn(attr), "Times New Roman")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def set_style_lang(style, lang="ru-RU"):
    rpr = style.element.get_or_add_rPr()
    existing = rpr.find(qn("w:lang"))
    if existing is not None:
        rpr.remove(existing)
    language = OxmlElement("w:lang")
    language.set(qn("w:val"), lang)
    language.set(qn("w:eastAsia"), lang)
    language.set(qn("w:bidi"), "ar-SA")
    rpr.append(language)


def make_style(doc, name, base="Normal"):
    styles = doc.styles
    if name in styles:
        return styles[name]
    style = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    style.base_style = styles[base]
    style.font.name = "Times New Roman"
    style.font.size = Pt(14)
    set_style_lang(style)
    return style


def set_outline_level(style, level):
    ppr = style.element.get_or_add_pPr()
    existing = ppr.find(qn("w:outlineLvl"))
    if existing is not None:
        ppr.remove(existing)
    outline = OxmlElement("w:outlineLvl")
    outline.set(qn("w:val"), str(level))
    ppr.append(outline)


def add_page_number_footer(section):
    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = Cm(0)
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)
    set_run_font(run, size=10)


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(3)
    section.right_margin = Cm(2)
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.footer_distance = Cm(1)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(14)
    set_style_lang(normal)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.first_line_indent = Cm(1.25)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    title = make_style(doc, "Report Title")
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.first_line_indent = Cm(0)
    title.paragraph_format.line_spacing = 1.15
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(0)

    meta = make_style(doc, "Report Meta")
    meta.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    meta.paragraph_format.first_line_indent = Cm(0)
    meta.paragraph_format.line_spacing = 1.5
    meta.paragraph_format.space_before = Pt(0)
    meta.paragraph_format.space_after = Pt(0)

    heading1 = make_style(doc, "Report Heading 1")
    heading1.font.bold = True
    heading1.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    heading1.paragraph_format.first_line_indent = Cm(1.25)
    heading1.paragraph_format.space_before = Pt(12)
    heading1.paragraph_format.space_after = Pt(6)
    heading1.paragraph_format.keep_with_next = True
    set_outline_level(heading1, 0)

    heading2 = make_style(doc, "Report Heading 2")
    heading2.font.bold = True
    heading2.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    heading2.paragraph_format.first_line_indent = Cm(1.25)
    heading2.paragraph_format.space_before = Pt(10)
    heading2.paragraph_format.space_after = Pt(4)
    heading2.paragraph_format.keep_with_next = True
    set_outline_level(heading2, 1)

    caption = make_style(doc, "Report Caption")
    caption.font.size = Pt(12)
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.first_line_indent = Cm(0)
    caption.paragraph_format.line_spacing = 1.0
    caption.paragraph_format.space_before = Pt(6)
    caption.paragraph_format.space_after = Pt(6)

    table_text = make_style(doc, "Report Table Text")
    table_text.font.size = Pt(10.5)
    table_text.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    table_text.paragraph_format.first_line_indent = Cm(0)
    table_text.paragraph_format.line_spacing = 1.0
    table_text.paragraph_format.space_before = Pt(0)
    table_text.paragraph_format.space_after = Pt(0)

    table_header = make_style(doc, "Report Table Header")
    table_header.font.size = Pt(10.5)
    table_header.font.bold = True
    table_header.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    table_header.paragraph_format.first_line_indent = Cm(0)
    table_header.paragraph_format.line_spacing = 1.0
    table_header.paragraph_format.space_before = Pt(0)
    table_header.paragraph_format.space_after = Pt(0)

    list_style = make_style(doc, "Report List Item")
    list_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    list_style.paragraph_format.left_indent = Cm(2.5)
    list_style.paragraph_format.first_line_indent = Cm(-0.8)
    list_style.paragraph_format.line_spacing = 1.5
    list_style.paragraph_format.space_before = Pt(0)
    list_style.paragraph_format.space_after = Pt(0)



def add_title_line(doc, text, bold=False):
    p = doc.add_paragraph(style="Report Title")
    run = p.add_run(text)
    set_run_font(run, size=14, bold=bold)
    return p


def add_meta_line(doc, text):
    p = doc.add_paragraph(style="Report Meta")
    run = p.add_run(text)
    set_run_font(run, size=14)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style="Report Heading 1" if level == 1 else "Report Heading 2")
    run = p.add_run(text)
    set_run_font(run, size=14, bold=True)
    return p


def add_body(doc, text, bold_prefix=None):
    p = doc.add_paragraph(style="Normal")
    if bold_prefix:
        run = p.add_run(bold_prefix)
        set_run_font(run, size=14, bold=True)
    run = p.add_run(text)
    set_run_font(run, size=14)
    return p


def add_label(doc, text):
    p = doc.add_paragraph(style="Normal")
    run = p.add_run(text)
    set_run_font(run, size=14, bold=True)
    return p


def add_numbered(doc, items, start=1):
    for index, text in enumerate(items, start=start):
        p = doc.add_paragraph(style="Report List Item")
        run = p.add_run(f"{index}.\t{text}")
        set_run_font(run, size=14)


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(table, top=40, start=80, bottom=40, end=80):
    tbl_pr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), str(value))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tbl_pr.append(mar)


def set_table_layout_fixed(table, widths_cm):
    tbl_pr = table._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tbl_pr.append(layout)
    tbl_w = OxmlElement("w:tblW")
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(int(sum(widths_cm) * 567)))
    tbl_pr.append(tbl_w)


def add_table(doc, headers, rows, widths_cm, first_column_center=True):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_layout_fixed(table, widths_cm)
    set_cell_margins(table)

    for index, width in enumerate(widths_cm):
        table.columns[index].width = Cm(width)

    header_cells = table.rows[0].cells
    for index, header in enumerate(headers):
        cell = header_cells[index]
        cell.width = Cm(widths_cm[index])
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        shade_cell(cell, "F2F2F2")
        paragraph = cell.paragraphs[0]
        paragraph.style = "Report Table Header"
        run = paragraph.add_run(header)
        set_run_font(run, size=10.5, bold=True)

    tr = table.rows[0]._tr
    tr_pr = tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tr_pr.append(tbl_header)

    for row_data in rows:
        row = table.add_row()
        for index, value in enumerate(row_data):
            cell = row.cells[index]
            cell.width = Cm(widths_cm[index])
            cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
            paragraph = cell.paragraphs[0]
            paragraph.style = "Report Table Text"
            if index == 0 and first_column_center:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = paragraph.add_run(str(value))
            set_run_font(run, size=10.5)
        row_pr = row._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        row_pr.append(cant_split)

    return table


def add_placeholder_table(doc):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_layout_fixed(table, [16.0])
    table.columns[0].width = Cm(16.0)
    row = table.rows[0]
    row.height = Cm(9.0)
    row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
    cell = row.cells[0]
    cell.width = Cm(16.0)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "dashed")
        el.set(qn("w:sz"), "6")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "BFBFBF")
        borders.append(el)
    tbl_pr.append(borders)

    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = Cm(0)
    run = paragraph.add_run("[ Место для рисунка 1 — вставьте ER-диаграмму ]")
    set_run_font(run, size=12, color=RGBColor(150, 150, 150))


# ---------------------------------------------------------------------------
# Документ
# ---------------------------------------------------------------------------

doc = Document()
configure_document(doc)

doc.core_properties.title = "УСР 1 — концептуальное проектирование базы данных VibeCity"
doc.core_properties.subject = "Системы баз данных"
doc.core_properties.author = ""
doc.core_properties.comments = "Отчёт по УСР 1 для проекта VibeCity (афиша мероприятий)."

add_title_line(doc, "Системы баз данных")
add_title_line(doc, "Групповой проект")
add_title_line(doc, "УСР 1 – концептуальное проектирование", bold=True)
doc.add_paragraph(style="Report Title")

add_meta_line(doc, "Предметная область: «Афиша мероприятий VibeCity»")
add_meta_line(doc, "Лидер команды: ______________________________")
add_meta_line(doc, "Участники команды: ______________________________")

doc.add_paragraph(style="Report Title")
add_heading(doc, "1.  Выбор предметной области и группы пользователей")

add_body(
    doc,
    " – информационная система афиши мероприятий и продажи билетов, предназначенная для "
    "учета пользователей, организаторов, площадок, мероприятий, билетов, отзывов, подарков "
    "и промокодов. В проекте рассматриваются публикация и модерация мероприятий, покупка "
    "билетов с выбором места, отзывы и управление справочниками. Платежный шлюз, "
    "видеотрансляции, чат, рекомендательная система и автоматический подбор мероприятий "
    "в модель не входят.",
    bold_prefix="Предметная область",
)

add_body(
    doc,
    " – гость, зритель (зарегистрированный пользователь), организатор и администратор. "
    "Гость просматривает афишу и карточки мероприятий. Зритель покупает билеты, хранит "
    "QR-коды, оставляет отзывы и редактирует профиль. Организатор создает и редактирует "
    "свои мероприятия, видит статистику продаж и отправляет карточки на модерацию. "
    "Администратор модерирует мероприятия, управляет пользователями и справочниками, "
    "удаляет отзывы и просматривает сводную статистику.",
    bold_prefix="Группы пользователей",
)

add_heading(doc, "2.  Неформальное описание предметной области")
add_heading(doc, "2.1.  Особенности учета информации", level=2)

add_numbered(doc, [
    "Каждому пользователю присваивается уникальный код; сохраняются фамилия, имя, отчество, "
    "телефон, email, хэш пароля, роль, признак блокировки и даты создания/обновления.",
    "Каждому организатору присваивается уникальный код; сохраняются название, контактное "
    "лицо и телефон.",
    "Каждому месту проведения присваивается уникальный код; сохраняются город, название, "
    "адрес и вместимость.",
    "Каждый тег имеет уникальный код, название и тип.",
    "Каждое мероприятие имеет уникальный код, название, дату, время, описание, категорию, "
    "ссылку на изображение, цену, статус, а также ссылки на организатора, место, тег и автора карточки.",
    "Один организатор может проводить много мероприятий, одно место может принимать много "
    "мероприятий, один тег может встречаться у многих мероприятий; у каждого мероприятия "
    "ровно один организатор, одно место и один тег.",
    "Статус мероприятия принимает одно из значений: «На модерации», «Опубликовано», "
    "«Отклонено»; гостям и зрителям видны только опубликованные мероприятия.",
    "Один пользователь может купить много билетов, одно мероприятие может иметь много "
    "билетов; одно и то же место на одно мероприятие не может быть продано дважды.",
    "Билет хранит номер места, цену на момент покупки, статус и уникальный QR-код; "
    "промокод для билета указывается необязательно.",
    "Промокод принадлежит подарку, имеет уникальный код, статус и скидку в процентах; "
    "после применения промокод становится использованным.",
    "Один пользователь может оставить только один отзыв на одно мероприятие; оценка "
    "принимает целое значение от 1 до 5, отзыв доступен только после покупки билета.",
    "Количество активных билетов на мероприятие не может превышать вместимость места проведения.",
    "Администратор может заблокировать пользователя; заблокированный пользователь не может "
    "войти в систему и купить билет.",
    "Любое редактирование мероприятия пользователем возвращает карточку на повторную модерацию.",
])

add_heading(doc, "2.2.  Информационные потребности пользователей", level=2)

add_label(doc, "Администратор:")
add_numbered(doc, [
    "Получать список мероприятий, ожидающих модерации, с автором, площадкой и датой.",
    "Одобрять или отклонять мероприятие и изменять его статус.",
    "Просматривать все мероприятия с фильтрами по статусу, категории и городу.",
    "Удалять мероприятия вместе со связанными билетами и отзывами.",
    "Искать пользователей и блокировать или разблокировать их.",
    "Вести справочники мест, организаторов, тегов, подарков и промокодов.",
    "Удалять отзывы и просматривать сводную статистику: пользователей, мероприятия, "
    "активные билеты, отзывы, выручку и количество карточек на модерации.",
])

add_label(doc, "Организатор:")
add_numbered(doc, [
    "Создавать мероприятие, выбирая организатора, место и тег из справочников.",
    "Просматривать список своих мероприятий и их статусы.",
    "Редактировать свои мероприятия и повторно отправлять их на модерацию.",
    "Видеть количество проданных билетов и оборот по своим мероприятиям.",
])

add_label(doc, "Зритель:")
add_numbered(doc, [
    "Просматривать афишу с фильтрами по категории, городу, тегу и датам, а также с поиском и пагинацией.",
    "Открывать карточку мероприятия с описанием, площадкой, организатором, отзывами и свободными местами.",
    "Покупать билет с выбором конкретного места и применением промокода.",
    "Просматривать свои билеты и показывать QR-код на входе.",
    "Оставлять и удалять свой отзыв после покупки билета, редактировать профиль.",
])

add_heading(doc, "3.  Выделение сущностей")

add_body(
    doc,
    "Сущности-объекты: «Пользователь», «Организатор», «Место проведения», «Тег», "
    "«Подарок», «Промокод». Сущности-процессы: «Мероприятие», «Билет», «Отзыв». "
    "Всего 9 сущностей, что определяется масштабом предметной области.",
)
add_body(
    doc,
    "Центральной сущностью является «Мероприятие»: она связывает организатора, место, тег "
    "и автора карточки, а от неё зависят билеты и отзывы. «Билет» связывает пользователя, "
    "мероприятие и (необязательно) промокод. «Промокод» связан с подарком и используется "
    "при покупке билета.",
)

add_heading(doc, "4.  Атрибуты, домены, ограничения и значения по умолчанию")
add_body(
    doc,
    "PK – первичный ключ; FK – внешний ключ; обязательность «Да» означает, что значение "
    "NULL не допускается. В концептуальной модели используются русские логические имена; "
    "в физической модели Prisma им соответствуют английские имена (User, Event, Ticket и т. д.).",
)

entity_tables = [
    ("4.1.  Пользователь", [
        ("PK", "КодПользователя", "Целое", "PRIMARY KEY, UNIQUE, NOT NULL, > 0", "Автоинкремент"),
        ("1", "Фамилия", "Строка(100)", "NOT NULL, непустая строка", "Не задано"),
        ("2", "Имя", "Строка(100)", "NOT NULL, непустая строка", "Не задано"),
        ("3", "Отчество", "Строка(100)", "NULL допускается", "NULL"),
        ("4", "Телефон", "Строка(20)", "NOT NULL, UNIQUE, формат +375 XX XXX-XX-XX", "Не задано"),
        ("5", "Email", "Строка(254)", "NOT NULL, UNIQUE, корректный формат", "Не задано"),
        ("6", "Пароль", "Строка(255)", "NOT NULL, bcrypt-хэш", "Не задано"),
        ("7", "Роль", "Перечисление", "NOT NULL, USER / ADMIN", "USER"),
        ("8", "ПризнакБлокировки", "Логическое", "NOT NULL", "false"),
        ("9", "ДатаСоздания", "Дата и время", "NOT NULL", "now()"),
        ("10", "ДатаОбновления", "Дата и время", "NOT NULL", "now()"),
    ]),
    ("4.2.  Организатор", [
        ("PK", "КодОрганизатора", "Целое", "PRIMARY KEY, UNIQUE, NOT NULL, > 0", "Автоинкремент"),
        ("1", "Название", "Строка(150)", "NOT NULL, непустая строка", "Не задано"),
        ("2", "КонтактноеЛицо", "Строка(150)", "NOT NULL, непустая строка", "Не задано"),
        ("3", "Телефон", "Строка(20)", "NOT NULL, непустая строка", "Не задано"),
    ]),
    ("4.3.  Место проведения", [
        ("PK", "КодМеста", "Целое", "PRIMARY KEY, UNIQUE, NOT NULL, > 0", "Автоинкремент"),
        ("1", "Город", "Строка(100)", "NOT NULL, непустая строка", "Не задано"),
        ("2", "Название", "Строка(150)", "NOT NULL, непустая строка", "Не задано"),
        ("3", "Адрес", "Строка(255)", "NOT NULL, непустая строка", "Не задано"),
        ("4", "Вместимость", "Целое", "NOT NULL, > 0", "Не задано"),
    ]),
    ("4.4.  Тег", [
        ("PK", "КодТега", "Целое", "PRIMARY KEY, UNIQUE, NOT NULL, > 0", "Автоинкремент"),
        ("1", "Название", "Строка(100)", "NOT NULL, непустая строка", "Не задано"),
        ("2", "Тип", "Строка(100)", "NOT NULL, непустая строка", "Не задано"),
    ]),
    ("4.5.  Мероприятие", [
        ("PK", "КодМероприятия", "Целое", "PRIMARY KEY, UNIQUE, NOT NULL, > 0", "Автоинкремент"),
        ("1", "КодОрганизатора", "Целое", "FK → Организатор, NOT NULL; существующий код", "Не задано"),
        ("2", "КодМеста", "Целое", "FK → Место проведения, NOT NULL; существующий код", "Не задано"),
        ("3", "КодТега", "Целое", "FK → Тег, NOT NULL; существующий код", "Не задано"),
        ("4", "КодАвтора", "Целое", "FK → Пользователь, NULL допускается; при удалении автора — NULL", "NULL"),
        ("5", "Название", "Строка(200)", "NOT NULL, непустая строка", "Не задано"),
        ("6", "Дата", "Дата", "NOT NULL, не ранее текущей даты при создании", "Не задано"),
        ("7", "Время", "Время", "NOT NULL, формат ЧЧ:ММ", "Не задано"),
        ("8", "Описание", "Текст", "NOT NULL, непустой текст", "Не задано"),
        ("9", "Категория", "Строка(100)", "NOT NULL, непустая строка", "Не задано"),
        ("10", "СсылкаНаФото", "Строка(2000)", "NULL допускается", "NULL"),
        ("11", "Цена", "Целое (р.)", "NOT NULL, >= 0", "0"),
        ("12", "Статус", "Перечисление", "NOT NULL, PENDING / APPROVED / REJECTED; для карточек организатора — PENDING", "APPROVED"),
        ("13", "ДатаСоздания", "Дата и время", "NOT NULL", "now()"),
        ("14", "ДатаОбновления", "Дата и время", "NOT NULL", "now()"),
    ]),
    ("4.6.  Отзыв", [
        ("PK", "КодОтзыва", "Целое", "PRIMARY KEY, UNIQUE, NOT NULL, > 0", "Автоинкремент"),
        ("1", "КодПользователя", "Целое", "FK → Пользователь, NOT NULL; существующий код", "Не задано"),
        ("2", "КодМероприятия", "Целое", "FK → Мероприятие, NOT NULL; существующий код", "Не задано"),
        ("3", "Текст", "Текст", "NOT NULL, непустой текст", "Не задано"),
        ("4", "Оценка", "Целое", "NOT NULL, от 1 до 5", "Не задано"),
        ("5", "Дата", "Дата и время", "NOT NULL", "now()"),
    ]),
    ("4.7.  Подарок", [
        ("PK", "КодПодарка", "Целое", "PRIMARY KEY, UNIQUE, NOT NULL, > 0", "Автоинкремент"),
        ("1", "Название", "Строка(150)", "NOT NULL, непустая строка", "Не задано"),
        ("2", "Описание", "Текст", "NOT NULL, непустой текст", "Не задано"),
        ("3", "Тип", "Строка(100)", "NOT NULL, непустая строка", "Не задано"),
        ("4", "Статус", "Перечисление", "NOT NULL, AVAILABLE / RESERVED / GIVEN", "AVAILABLE"),
    ]),
    ("4.8.  Промокод", [
        ("PK", "КодПромокода", "Целое", "PRIMARY KEY, UNIQUE, NOT NULL, > 0", "Автоинкремент"),
        ("1", "КодПодарка", "Целое", "FK → Подарок, NOT NULL; существующий код", "Не задано"),
        ("2", "Код", "Строка(50)", "NOT NULL, UNIQUE, непустая строка", "Не задано"),
        ("3", "Статус", "Перечисление", "NOT NULL, ACTIVE / USED / EXPIRED", "ACTIVE"),
        ("4", "Скидка", "Целое", "NOT NULL, от 0 до 100", "0"),
    ]),
    ("4.9.  Билет", [
        ("PK", "КодБилета", "Целое", "PRIMARY KEY, UNIQUE, NOT NULL, > 0", "Автоинкремент"),
        ("1", "КодПользователя", "Целое", "FK → Пользователь, NOT NULL; существующий код", "Не задано"),
        ("2", "КодМероприятия", "Целое", "FK → Мероприятие, NOT NULL; существующий код", "Не задано"),
        ("3", "КодПромокода", "Целое", "FK → Промокод, NULL допускается", "NULL"),
        ("4", "НомерМеста", "Целое", "NOT NULL, >= 1", "Не задано"),
        ("5", "Цена", "Целое (р.)", "NOT NULL, >= 0", "0"),
        ("6", "Статус", "Перечисление", "NOT NULL, ACTIVE / USED / CANCELLED", "ACTIVE"),
        ("7", "QR-код", "Строка(36)", "NOT NULL, UNIQUE, UUID", "Генерируется приложением"),
        ("8", "ДатаСоздания", "Дата и время", "NOT NULL", "now()"),
    ]),
]

entity_widths = [1.2, 4.0, 3.0, 5.0, 2.8]
for title, rows in entity_tables:
    add_heading(doc, title, level=2)
    add_table(
        doc,
        ["№", "Имя атрибута", "Домен", "Ограничения", "Значение по умолчанию"],
        rows,
        entity_widths,
    )
    if title.startswith("4.6."):
        add_body(doc, "Для сущности «Отзыв» действует составное ограничение UNIQUE (КодПользователя, КодМероприятия).")
    if title.startswith("4.9."):
        add_body(doc, "Для сущности «Билет» действует составное ограничение UNIQUE (КодМероприятия, НомерМеста).")

add_body(
    doc,
    "Дополнительные ограничения целостности: удаление организатора, места или тега при "
    "наличии связанных мероприятий запрещено до разрешения зависимостей; при удалении "
    "мероприятия каскадно удаляются его билеты и отзывы, при удалении пользователя — его "
    "билеты и отзывы, а созданные им мероприятия сохраняются с пустым автором. Промокод "
    "применяется одноразово, после покупки получает статус USED. Количество активных "
    "билетов не превышает вместимость места проведения; проверка выполняется логикой "
    "приложения. Цена, скидка и номер места неотрицательны, оценка находится в диапазоне "
    "1–5, а отзыв можно оставить только при наличии активного билета.",
)

add_heading(doc, "5.  Отношения между сущностями")

relation_widths = [4.4, 2.4, 9.2]
relation_rows = [
    ("Пользователь – Билет", "1:M", "У пользователя много билетов; каждый билет принадлежит одному пользователю."),
    ("Мероприятие – Билет", "1:M", "У мероприятия много билетов; каждый билет относится к одному мероприятию."),
    ("Пользователь – Отзыв", "1:M", "Пользователь может оставить много отзывов; каждый отзыв принадлежит одному пользователю."),
    ("Мероприятие – Отзыв", "1:M", "У мероприятия много отзывов; каждый отзыв относится к одному мероприятию."),
    ("Организатор – Мероприятие", "1:M", "У организатора много мероприятий; у мероприятия один организатор."),
    ("Место проведения – Мероприятие", "1:M", "На одной площадке проходит много мероприятий; у мероприятия одно место."),
    ("Тег – Мероприятие", "1:M", "Один тег может быть у многих мероприятий; у мероприятия один тег."),
    ("Пользователь – Мероприятие (автор)", "0..1:M", "Пользователь может создать много карточек; автор может быть не указан."),
    ("Подарок – Промокод", "1:M", "У подарка много промокодов; каждый промокод принадлежит одному подарку."),
    ("Промокод – Билет", "0..1:M", "Промокод может быть применён в билетах; у билета промокод может отсутствовать."),
]
add_table(doc, ["Связь", "Тип", "Описание"], relation_rows, relation_widths)

add_body(
    doc,
    "Связь «Пользователь – Мероприятие» через билеты и отзывы концептуально имеет тип M:M, "
    "но непосредственно не хранится: каждая пара реализуется через сущности-процессы "
    "«Билет» и «Отзыв». Аналогично связь «Подарок – Мероприятие» реализуется цепочкой "
    "«Подарок – Промокод – Билет – Мероприятие».",
)

add_heading(doc, "6.  Внешняя и концептуальная схемы")
add_body(
    doc,
    "В модели выделены три группы пользователей: зритель, организатор и администратор. "
    "Гость не имеет учетной записи и работает с публичной частью внешней схемы зрителя. "
    "Внешняя схема зрителя включает «Мероприятие», «Место проведения», «Организатор», "
    "«Тег», «Отзыв», «Билет», «Промокод» и «Подарок» (чтение), а также собственный профиль "
    "«Пользователь». Внешняя схема организатора включает «Мероприятие» (свои карточки), "
    "«Место проведения», «Организатор», «Тег», «Билет» (продажи по своим мероприятиям) и "
    "«Пользователь» (свой профиль). Внешняя схема администратора совпадает с "
    "концептуальной и включает все девять сущностей, поскольку администратор управляет "
    "пользователями, мероприятиями, отзывами и справочниками.",
)
add_body(
    doc,
    "Таким образом, внешние схемы групп являются подмножествами концептуальной схемы, а "
    "концептуальная схема описывает предметную область целиком.",
)
add_body(doc, "Ниже на рисунке 1 приведена ER-диаграмма в нотации «вороньей лапки»:")

doc.add_page_break()
add_placeholder_table(doc)
caption = doc.add_paragraph(style="Report Caption")
run = caption.add_run(
    "Рисунок 1. Концептуальная ER-диаграмма VibeCity.\n"
    "Девять сущностей и десять связей между ними."
)
set_run_font(run, size=12)

add_heading(doc, "7.  Итоги первого этапа")
add_body(
    doc,
    "Для предметной области «Афиша мероприятий VibeCity» выделены девять сущностей и "
    "десять связей типа 1:M, включая необязательные связи 0..1:M. Описаны атрибуты, "
    "домены, ограничения и значения по умолчанию, сформулированы информационные "
    "потребности трех групп пользователей, построена концептуальная ER-диаграмма. Модель "
    "может использоваться для последующей нормализации и физической реализации базы "
    "данных (MySQL 8.0, Prisma ORM).",
)

doc.save(OUT)
print("saved", OUT)
